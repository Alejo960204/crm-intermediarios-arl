import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from ..constantes import ROL_CLIENTE, ROL_EMPLEADO, ROL_ROOT
from ..repositories import usuarios_repository
from ..security import hash_password

MAX_BYTES = 2 * 1024 * 1024
MAX_FILAS = 500
MIN_CLAVE = 8
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_COLUMNAS = {  # nombre normalizado -> campo
    "nombre": "nombre", "nombres": "nombre",
    "correo": "correo", "correoelectronico": "correo", "email": "correo",
    "contrasena": "clave", "clave": "clave", "password": "clave",
    "rol": "rol", "perfil": "rol",
}


def _norm(t) -> str:
    t = unicodedata.normalize("NFD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", t.lower())


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # una contraseña numérica en Excel llega como número
    return str(v).strip()


def plantilla_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuarios"
    ws.append(["Nombre", "Correo", "Contraseña", "Rol"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    for col, ancho in zip("ABCD", (35, 35, 25, 15)):
        ws.column_dimensions[col].width = ancho
    info = wb.create_sheet("Instrucciones")
    for linea in [
        "Complete una fila por usuario en la hoja «Usuarios», empezando en la fila 2.",
        "Columnas obligatorias: Nombre, Correo y Contraseña (no cambie los títulos de la fila 1).",
        f"La contraseña debe tener al menos {MIN_CLAVE} caracteres.",
        "Rol (opcional): escriba «cliente» o «empleado». Si la deja vacía, el usuario se crea como «cliente».",
        "No se pueden crear administradores desde esta carga.",
        "Si el correo ya existe en el sistema, esa fila se omite (no se modifica el usuario).",
        f"Máximo {MAX_FILAS} filas por archivo.",
    ]:
        info.append([linea])
    info.column_dimensions["A"].width = 100
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _leer(contenido: bytes) -> list[tuple[int, dict]]:
    if not contenido.startswith(b"PK"):
        raise HTTPException(422, "El archivo no es un Excel válido (.xlsx).")
    try:
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
        filas = wb.worksheets[0].iter_rows(values_only=True)
        encabezado = next(filas, None)
    except Exception as e:  # archivo dañado o no soportado
        raise HTTPException(422, "No se pudo leer el Excel. Use la plantilla y guarde como .xlsx.") from e
    mapa = {}
    for i, titulo in enumerate(encabezado or []):
        campo = _COLUMNAS.get(_norm(titulo))
        if campo and campo not in mapa:
            mapa[campo] = i
    faltan = [c for c in ("nombre", "correo", "clave") if c not in mapa]
    if faltan:
        nombres = {"nombre": "Nombre", "correo": "Correo", "clave": "Contraseña"}
        raise HTTPException(422, "Faltan columnas en la fila 1: " + ", ".join(nombres[c] for c in faltan) + ".")
    datos = []
    for n, fila in enumerate(filas, start=2):
        vals = {c: _texto(fila[i]) if i < len(fila) else "" for c, i in mapa.items()}
        if not any(vals.values()):
            continue  # fila vacía
        datos.append((n, vals))
    if len(datos) > MAX_FILAS:
        raise HTTPException(422, f"El archivo tiene {len(datos)} filas; el máximo es {MAX_FILAS} por carga.")
    if not datos:
        raise HTTPException(422, "El archivo no tiene filas con datos.")
    return datos


def cargar_excel(contenido: bytes) -> dict:
    if not contenido:
        raise HTTPException(422, "El archivo está vacío.")
    if len(contenido) > MAX_BYTES:
        raise HTTPException(422, "El archivo supera el máximo de 2 MB.")
    datos = _leer(contenido)
    existentes = usuarios_repository.correos_existentes()
    vistos: dict[str, int] = {}
    resultado: dict[int, dict] = {}
    validos: list[tuple[int, str, str, str, str]] = []  # fila, nombre, correo, clave, rol

    for fila, v in datos:
        correo = v["correo"].lower()
        error = None
        rol = {"": ROL_CLIENTE, "cliente": ROL_CLIENTE, "empleado": ROL_EMPLEADO}.get(_norm(v.get("rol", "")))
        if rol is None:
            error = "Rol no válido (use «cliente» o «empleado», o déjelo vacío)"
        elif not v["nombre"]:
            error = "El nombre es obligatorio"
        elif len(v["nombre"]) > 150:
            error = "El nombre supera 150 caracteres"
        elif not correo or not _EMAIL_RE.match(correo) or len(correo) > 255:
            error = "Correo no válido"
        elif len(v["clave"]) < MIN_CLAVE:
            error = f"La contraseña debe tener al menos {MIN_CLAVE} caracteres"
        if error:
            resultado[fila] = {"fila": fila, "correo": v["correo"], "estado": "error", "detalle": error}
        elif correo in existentes:
            resultado[fila] = {"fila": fila, "correo": correo, "estado": "omitido", "detalle": "El correo ya existe"}
        elif correo in vistos:
            resultado[fila] = {"fila": fila, "correo": correo, "estado": "omitido",
                               "detalle": f"Correo repetido en el archivo (ya aparece en la fila {vistos[correo]})"}
        else:
            vistos[correo] = fila
            validos.append((fila, v["nombre"], correo, v["clave"], rol))

    # bcrypt es lento a propósito: se calcula en paralelo
    with ThreadPoolExecutor(max_workers=4) as pool:
        hashes = list(pool.map(hash_password, [x[3] for x in validos]))
    creados = usuarios_repository.crear_varios([
        {"nombre": n, "email": c, "password_hash": h, "rol": r}
        for (_, n, c, _, r), h in zip(validos, hashes)
    ])
    for fila, _, correo, _, rol in validos:
        ok = correo in creados
        resultado[fila] = {"fila": fila, "correo": correo, "estado": "creado" if ok else "omitido",
                           "detalle": f"Usuario creado con rol {rol}" if ok else "El correo ya existe"}

    filas = [resultado[k] for k in sorted(resultado)]
    return {
        "total": len(filas),
        "creados": sum(f["estado"] == "creado" for f in filas),
        "omitidos": sum(f["estado"] == "omitido" for f in filas),
        "con_error": sum(f["estado"] == "error" for f in filas),
        "filas": filas,
    }


# ---------- Req 14B: cambiar rol y activar / desactivar ----------
ROLES = (ROL_CLIENTE, ROL_EMPLEADO, ROL_ROOT)


def _existente(usuario_id: int) -> dict:
    u = usuarios_repository.obtener_con_datos(usuario_id)
    if not u:
        raise HTTPException(404, "Usuario no encontrado")
    return u


def cambiar_rol(actor: dict, usuario_id: int, rol: str) -> dict:
    if rol not in ROLES:
        raise HTTPException(422, detail=[{"loc": ["body", "rol"], "msg": "Rol no válido (cliente, empleado o root Admin)"}])
    if usuario_id == actor["id"]:
        raise HTTPException(403, "No puede cambiar su propio rol")
    u = _existente(usuario_id)
    if u["rol"] == rol:
        return u
    partes = []
    if u["empresas"]:
        partes.append(f"{u['empresas']} empresa(s)")
    if u["registros"]:
        partes.append(f"{u['registros']} registro(s)")
    if u["procesos_asignados"]:
        partes.append(f"{u['procesos_asignados']} proceso(s) asignado(s)")
    if partes:
        raise HTTPException(409, "No se puede cambiar el rol: el usuario tiene datos asociados ("
                            + ", ".join(partes) + "). Reasigne o retire esos datos primero.")
    usuarios_repository.cambiar_rol(usuario_id, rol)
    return _existente(usuario_id)


def cambiar_activo(actor: dict, usuario_id: int, activo: bool) -> dict:
    if usuario_id == actor["id"] and not activo:
        raise HTTPException(403, "No puede desactivar su propia cuenta")
    u = _existente(usuario_id)
    if u["activo"] != activo:
        usuarios_repository.cambiar_activo(usuario_id, activo)  # sus datos y asignaciones se conservan
    return _existente(usuario_id)
